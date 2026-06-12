"""
Metrobus Sunum Olusturucu
Calistir: python create_presentation.py
"""

import os
import tempfile
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Renk paleti ──────────────────────────────────────────────────
C_BG     = RGBColor(0x0F, 0x17, 0x2A)
C_ACCENT = RGBColor(0x00, 0xB4, 0xD8)
C_ACCENT2= RGBColor(0x90, 0xE0, 0xEF)
C_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
C_GRAY   = RGBColor(0xAA, 0xBB, 0xCC)
C_GREEN  = RGBColor(0x06, 0xD6, 0xA0)
C_YELLOW = RGBColor(0xFF, 0xD1, 0x66)
C_RED    = RGBColor(0xEF, 0x47, 0x6F)

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

_tmp_images = []   # temizlik için

# ── Yardımcılar ──────────────────────────────────────────────────

def new_prs():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color=C_BG):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color

def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(1, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    return s

def tb(slide, text, x, y, w, h, size=14, bold=False,
       color=C_WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf  = box.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    r   = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box

def header(slide, title, sub=None):
    tb(slide, title, Inches(0.7), Inches(0.2), Inches(11.9), Inches(0.7),
       size=28, bold=True, color=C_ACCENT)
    rect(slide, Inches(0.7), Inches(0.9), Inches(11.9), Inches(0.04), C_ACCENT)
    if sub:
        tb(slide, sub, Inches(0.7), Inches(0.94), Inches(11.9), Inches(0.35),
           size=12, color=C_GRAY)

def pgnum(slide, n):
    tb(slide, f"{n} / 10",
       Inches(11.9), Inches(7.05), Inches(1.2), Inches(0.35),
       size=11, color=C_GRAY, align=PP_ALIGN.RIGHT)

def formula_image(latex: str, width_in=5.5, height_in=0.8,
                  fontsize=18, bg_hex="#0F172A", fg_hex="#90E0EF") -> str:
    """LaTeX formülünü PNG'ye çevirir, yol döndürür."""
    fig, ax = plt.subplots(figsize=(width_in, height_in))
    fig.patch.set_facecolor(bg_hex)
    ax.set_facecolor(bg_hex)
    ax.axis("off")
    ax.text(0.5, 0.5, f"${latex}$",
            ha="center", va="center",
            fontsize=fontsize, color=fg_hex,
            transform=ax.transAxes,
            usetex=False)
    path = tempfile.mktemp(suffix=".png")
    plt.savefig(path, dpi=150, bbox_inches="tight",
                facecolor=bg_hex, edgecolor="none")
    plt.close(fig)
    _tmp_images.append(path)
    return path

def add_formula(slide, latex, x, y, w_in=5.5, h_in=0.75, fontsize=17):
    path = formula_image(latex, width_in=w_in, height_in=h_in, fontsize=fontsize)
    slide.shapes.add_picture(path, x, y, width=Inches(w_in), height=Inches(h_in))


# ════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ════════════════════════════════════════════════════════════════
def slide1(prs):
    sl = blank(prs); bg(sl)
    rect(sl, Inches(0), Inches(0), Inches(0.5), SLIDE_H, C_ACCENT)
    rect(sl, Inches(0.5), Inches(0), SLIDE_W - Inches(0.5), Inches(0.1), C_ACCENT)

    tb(sl, "Istanbul Metrobus",
       Inches(1.0), Inches(1.1), Inches(11), Inches(1.0),
       size=44, bold=True, color=C_WHITE)
    tb(sl, "Analytical Station-Slot Control System",
       Inches(1.0), Inches(2.05), Inches(11), Inches(0.7),
       size=26, color=C_ACCENT)
    rect(sl, Inches(1.0), Inches(2.8), Inches(5.5), Inches(0.05), C_ACCENT2)

    for i, (val, lbl) in enumerate([
        ("52 km", "Corridor"),
        ("45",    "Stations"),
        ("800K+", "Daily Passengers"),
        ("Real-Time", "Live Simulation"),
    ]):
        x = Inches(1.0) + Inches(i * 2.9)
        rect(sl, x, Inches(3.0), Inches(2.5), Inches(1.1),
             RGBColor(0x13, 0x22, 0x38), C_ACCENT)
        tb(sl, val,  x, Inches(3.05), Inches(2.5), Inches(0.6),
           size=22, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
        tb(sl, lbl,  x, Inches(3.65), Inches(2.5), Inches(0.4),
           size=11, color=C_GRAY, align=PP_ALIGN.CENTER)

    tb(sl, "Proactive, slot-based control  —  no AI, fully explainable",
       Inches(1.0), Inches(4.3), Inches(11), Inches(0.4),
       size=13, color=C_GRAY)
    pgnum(sl, 1)


# ════════════════════════════════════════════════════════════════
# SLIDE 2 — Problem & Objectives
# ════════════════════════════════════════════════════════════════
def slide2(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Problem Statement & Objectives")
    pgnum(sl, 2)

    rect(sl, Inches(0.7), Inches(1.1), Inches(5.6), Inches(5.6),
         RGBColor(0x18, 0x22, 0x38))
    tb(sl, "The Problem", Inches(0.9), Inches(1.2), Inches(5.0), Inches(0.45),
       size=15, bold=True, color=C_RED)
    for i, p in enumerate([
        "Vehicles cluster together (bus bunching)",
        "Large service gaps form on the route",
        "800,000+ daily passengers affected",
        "Traditional headway control reacts too late",
        "One overflowed station cascades delays",
    ]):
        tb(sl, f"  •  {p}", Inches(0.9), Inches(1.7) + Inches(i*0.72),
           Inches(5.2), Inches(0.65), size=13, color=C_WHITE)

    rect(sl, Inches(6.7), Inches(1.1), Inches(5.9), Inches(5.6),
         RGBColor(0x0D, 0x20, 0x2E))
    tb(sl, "Objectives", Inches(6.9), Inches(1.2), Inches(5.5), Inches(0.45),
       size=15, bold=True, color=C_GREEN)
    for i, o in enumerate([
        "Prevent station slot overflow",
        "Minimize cascade queue effects",
        "Real-time speed recommendations",
        "Purely analytical — no ML/RL",
        "Scalable: O(S×2) per tick",
    ]):
        tb(sl, f"  •  {o}", Inches(6.9), Inches(1.7) + Inches(i*0.72),
           Inches(5.5), Inches(0.65), size=13, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 3 — Background & Motivation
# ════════════════════════════════════════════════════════════════
def slide3(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Background & Motivation")
    pgnum(sl, 3)

    for i, (val, lbl, color) in enumerate([
        ("52 km",  "Corridor Length",    C_ACCENT),
        ("45",     "Stations",           C_GREEN),
        ("800K+",  "Daily Passengers",   C_YELLOW),
        ("~2 min", "Target Headway",     C_ACCENT2),
    ]):
        x = Inches(0.6) + Inches(i * 3.1)
        rect(sl, x, Inches(1.1), Inches(2.8), Inches(1.7),
             RGBColor(0x13, 0x1F, 0x35))
        tb(sl, val, x, Inches(1.2), Inches(2.8), Inches(0.85),
           size=32, bold=True, color=color, align=PP_ALIGN.CENTER)
        tb(sl, lbl, x, Inches(2.05), Inches(2.8), Inches(0.5),
           size=12, color=C_GRAY, align=PP_ALIGN.CENTER)

    tb(sl, "Core Insight",
       Inches(0.7), Inches(3.1), Inches(11.9), Inches(0.4),
       size=15, bold=True, color=C_ACCENT)
    rect(sl, Inches(0.7), Inches(3.5), Inches(11.9), Inches(0.04), C_ACCENT)

    insights = [
        ("Root Cause",     "Station slot overflow — not headway deviation"),
        ("Old Approach",   "Measure headway gap → react with hold/skip → too late"),
        ("New Approach",   "Predict slot availability → adjust speed before arriving"),
        ("Key Property",   "Proactive control eliminates the need for reactive correction"),
    ]
    for i, (term, desc) in enumerate(insights):
        y = Inches(3.65) + Inches(i * 0.78)
        tb(sl, term, Inches(0.9), y, Inches(2.8), Inches(0.65),
           size=13, bold=True, color=C_ACCENT2)
        tb(sl, desc, Inches(3.7), y, Inches(8.8), Inches(0.65),
           size=13, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 4 — Design Methodology
# ════════════════════════════════════════════════════════════════
def slide4(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Design Methodology")
    pgnum(sl, 4)

    # Pipeline
    for i, (lbl, color) in enumerate([
        ("Physics\nEngine",              C_GRAY),
        ("Smart Stops\nx45",             C_ACCENT),
        ("Stop\nCoordination",           C_ACCENT2),
        ("Speed\nCommands",              C_GREEN),
        ("Live\nDashboard",              C_YELLOW),
    ]):
        x = Inches(0.5) + Inches(i * 2.55)
        rect(sl, x, Inches(1.1), Inches(2.15), Inches(1.4),
             RGBColor(0x12, 0x20, 0x36), color)
        tb(sl, lbl, x, Inches(1.15), Inches(2.15), Inches(1.3),
           size=12, bold=True, color=color, align=PP_ALIGN.CENTER)
        if i < 4:
            tb(sl, ">", x + Inches(2.18), Inches(1.65),
               Inches(0.3), Inches(0.4),
               size=16, bold=True, color=C_GRAY, align=PP_ALIGN.CENTER)

    tb(sl, "Design Principles",
       Inches(0.7), Inches(2.75), Inches(11.9), Inches(0.4),
       size=15, bold=True, color=C_ACCENT)
    rect(sl, Inches(0.7), Inches(3.15), Inches(11.9), Inches(0.04), C_ACCENT)

    principles = [
        (C_ACCENT,  "Purely Analytical",   "Deterministic kinematic formulas — no ML, no training data"),
        (C_GREEN,   "Decentralized",        "Each stop independently manages its own zone (~1-2 vehicles)"),
        (C_YELLOW,  "Single Thread",        "asyncio cooperative multitasking — simple, no race conditions"),
        (C_ACCENT2, "Forward Safety",       "Collision prevention layer always overrides speed commands"),
    ]
    for i, (color, term, desc) in enumerate(principles):
        col = i % 2
        row = i // 2
        x = Inches(0.6) + Inches(col * 6.4)
        y = Inches(3.3) + Inches(row * 1.6)
        rect(sl, x, y, Inches(6.1), Inches(1.4), RGBColor(0x10, 0x1C, 0x2E))
        rect(sl, x, y, Inches(0.1), Inches(1.4), color)
        tb(sl, term, x + Inches(0.2), y + Inches(0.1), Inches(5.8), Inches(0.4),
           size=13, bold=True, color=color)
        tb(sl, desc, x + Inches(0.2), y + Inches(0.55), Inches(5.8), Inches(0.7),
           size=12, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 5 — Implementation I
# ════════════════════════════════════════════════════════════════
def slide5(prs):
    sl = blank(prs); bg(sl)
    header(sl, "How It Works I — Vehicle Movement & Stations")
    pgnum(sl, 5)

    # FSM kutuları
    phases = ["cruising", "approaching", "docking", "stopped", "doorsClosed", "departing"]
    colors = [C_GRAY, C_ACCENT, C_ACCENT2, C_GREEN, C_YELLOW, C_GRAY]
    for i, (phase, color) in enumerate(zip(phases, colors)):
        x = Inches(0.55) + Inches(i * 2.13)
        rect(sl, x, Inches(1.2), Inches(1.95), Inches(0.6),
             RGBColor(0x10, 0x1E, 0x32), color)
        tb(sl, phase, x, Inches(1.27), Inches(1.95), Inches(0.45),
           size=11, bold=True, color=color, align=PP_ALIGN.CENTER)
        if i < 5:
            tb(sl, ">", x + Inches(1.98), Inches(1.38),
               Inches(0.15), Inches(0.3),
               size=13, color=C_GRAY, align=PP_ALIGN.CENTER)

    tb(sl, "Vehicle Phase State Machine",
       Inches(0.6), Inches(1.05), Inches(8), Inches(0.2),
       size=11, color=C_GRAY)

    # İki kolon: sol parametreler, sağ SmartStop zone
    rect(sl, Inches(0.6), Inches(2.0), Inches(5.8), Inches(4.9),
         RGBColor(0x10, 0x1C, 0x2E))
    tb(sl, "Realistic Movement", Inches(0.8), Inches(2.1), Inches(5.4), Inches(0.4),
       size=14, bold=True, color=C_ACCENT)

    moves = [
        "Vehicles move with real-time physics",
        "Smooth, comfortable acceleration & braking",
        "Each vehicle safely follows the one ahead",
        "Stops have a limited number of platform slots",
        "Buses enter the platform from the rear",
        "Full buses form a queue outside the platform",
    ]
    for i, m in enumerate(moves):
        y = Inches(2.6) + Inches(i * 0.68)
        tb(sl, f"  •  {m}", Inches(0.85), y, Inches(5.4), Inches(0.6),
           size=12, color=C_WHITE)

    rect(sl, Inches(6.8), Inches(2.0), Inches(6.0), Inches(4.9),
         RGBColor(0x0D, 0x1E, 0x2E))
    tb(sl, "SmartStop Zone Logic", Inches(7.0), Inches(2.1), Inches(5.6), Inches(0.4),
       size=14, bold=True, color=C_ACCENT)

    zone_points = [
        "Each stop owns a zone [prev_stop → this_stop]",
        "Filters only vehicles in its own zone (~1-2)",
        "Computes ETA for each approaching vehicle",
        "Checks slot timeline for earliest free slot",
        "Runs cost-benefit: slow down or accept queue?",
        "Publishes zone state to StopInterface",
    ]
    for i, p in enumerate(zone_points):
        tb(sl, f"  {i+1}.  {p}",
           Inches(7.0), Inches(2.6) + Inches(i * 0.65),
           Inches(5.7), Inches(0.6),
           size=12, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 6 — Implementation II: Mathematics
# ════════════════════════════════════════════════════════════════
def slide6(prs):
    sl = blank(prs); bg(sl)
    header(sl, "How It Works II — How a Stop Makes a Decision")
    pgnum(sl, 6)

    steps = [
        (C_ACCENT,  "1.  Who is coming?",
         "Each stop watches only the few vehicles inside its own zone."),
        (C_ACCENT2, "2.  When will a slot free up?",
         "It estimates how long the buses now at the platform will stay."),
        (C_YELLOW,  "3.  Will this bus have to wait?",
         "It predicts whether the bus arrives to a free slot — or to the back of a queue."),
        (C_GREEN,   "4.  Is slowing down worth it?",
         "It weighs the queue and the knock-on delays against a small slowdown."),
        (C_RED,     "5.  Give the advice",
         "If worthwhile, it suggests a speed so the bus arrives just as a slot opens — no stop-and-go."),
    ]
    for i, (color, title, desc) in enumerate(steps):
        y = Inches(1.25) + Inches(i * 1.12)
        rect(sl, Inches(0.7), y, Inches(11.9), Inches(0.95),
             RGBColor(0x10, 0x1C, 0x2E))
        rect(sl, Inches(0.7), y, Inches(0.12), Inches(0.95), color)
        tb(sl, title, Inches(1.0), y + Inches(0.12), Inches(4.6), Inches(0.7),
           size=15, bold=True, color=color)
        tb(sl, desc, Inches(5.5), y + Inches(0.12), Inches(6.9), Inches(0.75),
           size=13, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 7 — Implementation III: Driver Speed-Band Advisory
# ════════════════════════════════════════════════════════════════
def slide7(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Driver Speed-Band Advisory")
    pgnum(sl, 7)

    # Üst kavram bandı
    rect(sl, Inches(0.6), Inches(1.05), Inches(12.1), Inches(0.95),
         RGBColor(0x0A, 0x1E, 0x30))
    rect(sl, Inches(0.6), Inches(1.05), Inches(0.1), Inches(0.95), C_ACCENT)
    tb(sl, "Advise a holdable BAND, not a point speed",
       Inches(0.85), Inches(1.12), Inches(11.5), Inches(0.4),
       size=15, bold=True, color=C_ACCENT)
    tb(sl, "A real driver cannot track \"go 61 km/h, now 62, now 60\" each second. "
           "We tell the driver \"keep 60-70 km/h\" — analogous to GLOSA speed advisory at traffic lights.",
       Inches(0.85), Inches(1.52), Inches(11.6), Inches(0.45),
       size=11.5, color=C_GRAY)

    # Sol kutu — neden band
    rect(sl, Inches(0.6), Inches(2.15), Inches(5.85), Inches(2.55),
         RGBColor(0x10, 0x1C, 0x2E))
    tb(sl, "Why a Band, Not a Number", Inches(0.8), Inches(2.25), Inches(5.5), Inches(0.4),
       size=13, bold=True, color=C_GREEN)
    for i, (k, v) in enumerate([
        ("Smoothing", "ignores second-to-second noise in the target"),
        ("Hold", "keeps a band steady instead of changing constantly"),
        ("Escape", "updates early only when the situation truly changes"),
    ]):
        y = Inches(2.7) + Inches(i * 0.62)
        tb(sl, f"  •  {k}", Inches(0.8), y, Inches(2.4), Inches(0.55),
           size=12, bold=True, color=C_ACCENT2)
        tb(sl, v, Inches(0.95), y + Inches(0.27), Inches(5.4), Inches(0.35),
           size=10.5, color=C_WHITE)

    # Sağ kutu — sonuç / metrik
    rect(sl, Inches(6.85), Inches(2.15), Inches(5.85), Inches(2.55),
         RGBColor(0x06, 0x2A, 0x1A))
    tb(sl, "Result", Inches(7.05), Inches(2.25), Inches(5.5), Inches(0.4),
       size=13, bold=True, color=C_GREEN)
    tb(sl, "171  →  2", Inches(7.05), Inches(2.65), Inches(5.5), Inches(0.9),
       size=46, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
    tb(sl, "band changes per 60 s of simulation",
       Inches(7.05), Inches(3.55), Inches(5.5), Inches(0.4),
       size=12, color=C_GRAY, align=PP_ALIGN.CENTER)
    tb(sl, "No more per-second \"speed up / slow down\" flicker",
       Inches(7.05), Inches(3.95), Inches(5.5), Inches(0.5),
       size=11, color=C_WHITE, align=PP_ALIGN.CENTER)

    # Alt — yaşam döngüsü
    tb(sl, "Advisory Lifecycle",
       Inches(0.6), Inches(4.85), Inches(11.9), Inches(0.35),
       size=13, bold=True, color=C_ACCENT)
    rect(sl, Inches(0.6), Inches(5.2), Inches(12.1), Inches(0.03), C_ACCENT)
    stages = [
        (C_YELLOW, "SLOW DOWN",
         "early arrival risks slot overflow → band drops (e.g. 50-60)"),
        (C_ACCENT, "HOLD",
         "driver keeps the band; metering happens while cruising"),
        (C_GREEN,  "RESUME",
         "risk cleared → band climbs back to full speed (cue to speed up)"),
    ]
    for i, (color, title, desc) in enumerate(stages):
        x = Inches(0.6) + Inches(i * 4.1)
        rect(sl, x, Inches(5.35), Inches(3.8), Inches(1.5),
             RGBColor(0x10, 0x1C, 0x2E))
        rect(sl, x, Inches(5.35), Inches(3.8), Inches(0.06), color)
        tb(sl, title, x + Inches(0.15), Inches(5.45), Inches(3.5), Inches(0.4),
           size=14, bold=True, color=color)
        tb(sl, desc, x + Inches(0.15), Inches(5.9), Inches(3.5), Inches(0.9),
           size=11, color=C_WHITE)
        if i < 2:
            tb(sl, ">", x + Inches(3.82), Inches(5.85), Inches(0.28), Inches(0.4),
               size=18, bold=True, color=C_GRAY, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════
# SLIDE 8 — Challenges & Solutions
# ════════════════════════════════════════════════════════════════
def slide8(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Challenges & Solutions")
    pgnum(sl, 8)

    # Üst bant — ML'den analitik sisteme geçiş
    rect(sl, Inches(0.5), Inches(1.1), Inches(12.3), Inches(1.65),
         RGBColor(0x0A, 0x1E, 0x2E))
    rect(sl, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.06), C_RED)
    tb(sl, "Initial Approach: AI / Reinforcement Learning",
       Inches(0.7), Inches(1.18), Inches(8), Inches(0.4),
       size=14, bold=True, color=C_RED)
    tb(sl, "The system was initially designed using a reinforcement learning agent to learn optimal speed policies. "
           "However, training instability, lack of interpretability, and the need for large amounts of simulation data "
           "made this approach impractical for a real-time transit control system.",
       Inches(0.7), Inches(1.62), Inches(8.5), Inches(0.95),
       size=11, color=C_GRAY)

    # Ok + sonuç kutusu
    tb(sl, "-->", Inches(9.2), Inches(1.75), Inches(0.6), Inches(0.5),
       size=20, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)
    rect(sl, Inches(9.9), Inches(1.55), Inches(2.8), Inches(1.1),
         RGBColor(0x06, 0x2A, 0x1A))
    rect(sl, Inches(9.9), Inches(1.55), Inches(2.8), Inches(0.06), C_GREEN)
    tb(sl, "Switched to\nAnalytical Control",
       Inches(9.95), Inches(1.62), Inches(2.7), Inches(0.55),
       size=12, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER)
    tb(sl, "Deterministic · Interpretable\nNo training required",
       Inches(9.95), Inches(2.17), Inches(2.7), Inches(0.4),
       size=10, color=C_ACCENT2, align=PP_ALIGN.CENTER)

    # Alt dört teknik zorluk
    items = [
        (C_YELLOW, "Reaction Time",
         "A speed advice reaches the driver a few seconds later",
         "We plan ahead for that delay before sending the advice"),
        (C_ACCENT, "Slot Reservation Chain",
         "Several buses may aim for the same free slot",
         "Each bus reserves its slot, so the next one sees it as taken"),
        (C_RED,    "Physical Slot Access",
         "Empty front slots are unreachable — buses enter from the rear",
         "We only count slots actually reachable from the back"),
        (C_ACCENT2,"Realistic Stop Times",
         "A fixed stop time ignored real-world variation",
         "We use field-measured stop times that vary trip to trip"),
    ]

    for i, (color, title, problem, solution) in enumerate(items):
        col = i % 2
        row = i // 2
        x = Inches(0.5) + Inches(col * 6.4)
        y = Inches(2.95) + Inches(row * 2.2)
        rect(sl, x, y, Inches(6.1), Inches(2.0), RGBColor(0x10, 0x1C, 0x2E))
        rect(sl, x, y, Inches(6.1), Inches(0.05), color)
        tb(sl, title,
           x + Inches(0.15), y + Inches(0.1), Inches(5.8), Inches(0.38),
           size=13, bold=True, color=color)
        tb(sl, "Challenge:  " + problem,
           x + Inches(0.15), y + Inches(0.52), Inches(5.8), Inches(0.5),
           size=11, color=C_GRAY)
        tb(sl, "Solution:  " + solution,
           x + Inches(0.15), y + Inches(1.07), Inches(5.8), Inches(0.82),
           size=11, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 9 — Conclusion
# ════════════════════════════════════════════════════════════════
def slide9(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Conclusion")
    pgnum(sl, 9)

    points = [
        (C_ACCENT,  "Problem Reframed",
         "Bus bunching is a station slot overflow problem — not a headway problem. "
         "Proactive slot-based control prevents bunching before it occurs."),
        (C_GREEN,   "Analytical & Interpretable",
         "Every decision follows deterministic kinematic formulas. "
         "No black-box models, no training, no data dependency."),
        (C_YELLOW,  "Efficient by Design",
         "Work stays light even with a large fleet — each stop only looks at "
         "its own handful of vehicles."),
        (C_ACCENT2, "Real-Time Capable",
         "Runs live with a real-time dashboard, on a single machine — "
         "simple to deploy, no special infrastructure needed."),
        (C_RED,     "Extensible",
         "Built from independent parts that can each be tested and improved "
         "on their own."),
    ]

    for i, (color, title, body) in enumerate(points):
        y = Inches(1.1) + Inches(i * 1.15)
        rect(sl, Inches(0.6), y, Inches(0.12), Inches(0.9), color)
        tb(sl, title, Inches(0.9), y, Inches(11.5), Inches(0.38),
           size=14, bold=True, color=color)
        tb(sl, body, Inches(0.9), y + Inches(0.4), Inches(11.5), Inches(0.65),
           size=12, color=C_WHITE)


# ════════════════════════════════════════════════════════════════
# SLIDE 10 — Future Work & Q&A
# ════════════════════════════════════════════════════════════════
def slide10(prs):
    sl = blank(prs); bg(sl)
    header(sl, "Future Work & Q&A")
    pgnum(sl, 10)

    # Ana madde — Yolcu metrikleri (büyük kutu)
    rect(sl, Inches(0.6), Inches(1.1), Inches(12.1), Inches(3.2),
         RGBColor(0x0A, 0x1E, 0x30))
    rect(sl, Inches(0.6), Inches(1.1), Inches(0.12), Inches(3.2), C_RED)
    tb(sl, "Passenger Density Metrics",
       Inches(0.9), Inches(1.18), Inches(11.5), Inches(0.45),
       size=18, bold=True, color=C_RED)
    tb(sl, "The current system has no awareness of passenger load. "
           "Integrating real-time occupancy data — both inside vehicles and on station platforms — "
           "would enable demand-aware slot scheduling and more accurate dwell time estimation.",
       Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.85),
       size=13, color=C_WHITE)
    for i, (lbl, desc) in enumerate([
        ("In-vehicle load",    "Adjust dwell time and slot priority based on passenger count"),
        ("Platform occupancy", "Detect crowded platforms early — trigger upstream slowdowns"),
        ("Demand prediction",  "Historical patterns could improve ideal arrival time estimation"),
    ]):
        y = Inches(2.6) + Inches(i * 0.55)
        tb(sl, f"  •  {lbl}:", Inches(0.9), y, Inches(2.8), Inches(0.48),
           size=12, bold=True, color=C_ACCENT2)
        tb(sl, desc, Inches(3.7), y, Inches(8.8), Inches(0.48),
           size=12, color=C_GRAY)

    # İki yan madde
    for i, (color, title, body) in enumerate([
        (C_ACCENT, "Smarter Cascade Model",
         "Account for delay that builds up across a whole chain of following "
         "buses — not just the next one — for more accurate decisions."),
        (C_GREEN,  "Higher-Fidelity Movement",
         "A more realistic acceleration, comfort and traffic model would "
         "sharpen both the speed advice and the decision engine further."),
    ]):
        x = Inches(0.6) + Inches(i * 6.4)
        rect(sl, x, Inches(4.5), Inches(6.1), Inches(2.3),
             RGBColor(0x10, 0x1C, 0x2E))
        rect(sl, x, Inches(4.5), Inches(0.1), Inches(2.3), color)
        tb(sl, title, x + Inches(0.2), Inches(4.6), Inches(5.8), Inches(0.4),
           size=13, bold=True, color=color)
        tb(sl, body, x + Inches(0.2), Inches(5.05), Inches(5.8), Inches(1.6),
           size=11, color=C_WHITE)

    # Q&A
    rect(sl, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35),
         RGBColor(0x00, 0x3A, 0x55))
    tb(sl, "Questions & Discussion",
       Inches(0.6), Inches(7.02), Inches(12.1), Inches(0.3),
       size=13, bold=True, color=C_ACCENT, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════════
# MAIN
# ════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    prs = new_prs()
    slide1(prs); slide2(prs); slide3(prs); slide4(prs); slide5(prs)
    slide6(prs); slide7(prs); slide8(prs); slide9(prs); slide10(prs)

    out = "Metrobus_Presentation_v2.pptx"
    prs.save(out)
    print(f"Kaydedildi: {out}")

    for p in _tmp_images:
        try:
            os.remove(p)
        except Exception:
            pass
